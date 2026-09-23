"""Flatten scraped journal JSON records into a single Excel workbook for easy review.

Reads every data_scrapes/json/<issn>.json and writes one row per (journal,
article type) to data_scrapes/journals_overview.xlsx. Source URLs are joined
into a single "; "-separated cell rather than spread across columns, since
the number of source pages per article type varies. Any missing value is
written as "NA" in the spreadsheet -- the underlying JSON keeps proper
`null`/`[]` so downstream code can still tell "missing" apart from the
literal string "NA".

Run directly: `python journal_scrapers/common/json_to_excel.py`
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

DATA_SCRAPES_DIR = Path(__file__).resolve().parents[2] / "data_scrapes"
OUTPUT_PATH = DATA_SCRAPES_DIR / "journals_overview.xlsx"

NA = "NA"

COLUMNS = [
    "Journal", "Publisher", "ISSN (print)", "ISSN (electronic)", "LLM", "Validated", "Date scraped",
    "Article type", "Total word limit (min)", "Total word limit (max)", "Total word limit (unit)",
    "Total word limit (excludes)", "Total word limit (notes)", "Description",
    "Structure", "Figures/tables", "Sources of info",
]


def _na(value):
    """Blank/None/empty collection -> "NA"; everything else passed through as-is."""
    if value is None:
        return NA
    if isinstance(value, (list, tuple)) and not value:
        return NA
    if isinstance(value, str) and not value.strip():
        return NA
    return value


def _join(items) -> str:
    return "; ".join(str(i) for i in items) if items else ""


def record_to_rows(record: dict) -> list[list]:
    """One row per article type; journal-level fields repeat across those rows."""
    issn = record.get("issn") or {}
    article_types = record.get("article_types") or [{}]

    rows = []
    for at in article_types:
        word_limit = at.get("total_word_limit") or {}
        rows.append([_na(v) for v in [
            record.get("journal"),
            record.get("publisher"),
            issn.get("print"),
            issn.get("electronic"),
            record.get("LLM"),
            record.get("validated"),
            record.get("date_scraped"),
            at.get("type"),
            word_limit.get("min"),
            word_limit.get("max"),
            word_limit.get("unit"),
            _join(word_limit.get("excludes")),
            word_limit.get("notes"),
            at.get("description"),
            at.get("structure"),
            at.get("figures_tables"),
            _join(at.get("source_urls")),
        ]])
    return rows


def build_workbook(records: list[dict]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Journals"
    ws.append(COLUMNS)
    ws.freeze_panes = "A2"

    for record in records:
        for row in record_to_rows(record):
            ws.append(row)

    for idx, header in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = min(max(len(header) + 2, 14), 45)

    return wb


def load_latest_records(data_dir: Path = DATA_SCRAPES_DIR) -> list[dict]:
    records = []
    for record_path in sorted((data_dir / "json").glob("*.json")):
        with record_path.open(encoding="utf-8") as f:
            records.append(json.load(f))
    return records


def main() -> Path:
    records = load_latest_records()
    if not records:
        raise SystemExit(f"No *.json files found under {DATA_SCRAPES_DIR / 'json'}")
    wb = build_workbook(records)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = main()
    print(f"Wrote {path}")
